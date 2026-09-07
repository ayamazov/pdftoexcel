#!/usr/bin/env python3
"""Build the ricostruzione workbook (fogli 'originale' + 'data valuta') from rows.json.

    python build_workbook.py rows-1993.json rows-1994.json --out ricostruzione.xlsx --verifica

One workbook for ALL years. Sign convention: saldo = F(prev) + dare - avere,
so positive saldo = conto a debito. Read references/output-format.md before
changing anything here, and remember a client-supplied target workbook wins.
"""
import argparse, glob, json
from datetime import datetime
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

HEAD = [" operazione", "valuta", "descrizione operazione", "dare", "avere", "saldo"]
WIDTH = {"A": 11.57, "B": 12.86, "C": 50.86, "D": 12.71, "F": 13.57}
PAD = 200  # blank formula rows so next year is a paste into A:E


def d(s):
    return datetime.strptime(s, "%Y-%m-%d")


def main():
    p = argparse.ArgumentParser()
    p.add_argument("files", nargs="+")
    p.add_argument("--out", required=True)
    p.add_argument("--verifica", action="store_true")
    a = p.parse_args()

    sts = []
    for pat in a.files:
        for f in sorted(glob.glob(pat)) or [pat]:
            j = json.load(open(f))
            for s in j["statements"]:
                s.setdefault("currency", j.get("currency", "ITL"))
                sts.append(s)
    sts.sort(key=lambda s: s["estratto_al"])

    cur = sts[0]["currency"]
    numfmt = "#,##0.00" if cur == "EUR" else "#,##0"
    datefmt = "D/M/YYYY"

    rows = []
    for s in sts:
        for x in s["rows"]:
            rows.append(dict(x, _st=s["estratto_al"]))

    opening = sts[0]
    wb = Workbook()

    def sheet(ws, hdr_row, data):
        if hdr_row == 2:
            ws.append([])
        ws.append(HEAD)
        first = hdr_row + 1
        # opening balance row: both dates, no descrizione, hardcoded saldo
        op0 = opening.get("saldo_iniziale_al", opening["estratto_al"])
        ws.cell(first, 1, d(op0)).number_format = datefmt
        ws.cell(first, 2, d(op0)).number_format = datefmt
        c = ws.cell(first, 6, opening["saldo_iniziale"])
        c.number_format = numfmt
        c.comment = Comment("SALDO INIZIALE, estratto al %s, pag. PDF %s"
                            % (op0,
                               ",".join(str(v) for v in opening.get("pdf_pages", []))), "skill")
        r = first + 1
        for x in data:
            ws.cell(r, 1, d(x["op"])).number_format = datefmt
            ws.cell(r, 2, d(x["val"])).number_format = datefmt
            ws.cell(r, 3, x["desc"])
            if x.get("dare") is not None:
                ws.cell(r, 4, x["dare"]).number_format = numfmt
            if x.get("avere") is not None:
                ws.cell(r, 5, x["avere"]).number_format = numfmt
            f = ws.cell(r, 6, "=F%d+D%d-E%d" % (r - 1, r, r))
            f.number_format = numfmt
            if x.get("raw") or x.get("page"):
                f.comment = Comment("%s  (pag. PDF %s)" % (x.get("raw", ""), x.get("page", "?")),
                                    "skill")
            r += 1
        for k in range(r, r + PAD):                       # room for next year
            ws.cell(k, 6, "=F%d+D%d-E%d" % (k - 1, k, k)).number_format = numfmt
        for col, w in WIDTH.items():
            ws.column_dimensions[col].width = w
        return r - 1

    ws1 = wb.active
    ws1.title = "originale"
    last1 = sheet(ws1, 2, sorted(rows, key=lambda x: x["op"]))

    ws2 = wb.create_sheet("data valuta")
    byval = sorted(range(len(rows)), key=lambda i: (rows[i]["val"], i))   # stable
    last2 = sheet(ws2, 1, [rows[i] for i in byval])
    ws2.freeze_panes = "B2"

    if a.verifica:
        v = wb.create_sheet("Verifica")
        v.append(["estratto al", "saldo calcolato", "saldo stampato", "differenza", "esito"])
        run = opening["saldo_iniziale"]
        for s in sts:
            run = s["saldo_iniziale"] + sum((x.get("dare") or 0) - (x.get("avere") or 0)
                                            for x in s["rows"])
            n = v.max_row + 1
            v.append([d(s["estratto_al"]), run, s["saldo_finale"], run - s["saldo_finale"]])
            v.cell(n, 1).number_format = datefmt
            for col in (2, 3, 4):
                v.cell(n, col).number_format = numfmt
            v.cell(n, 5, '=IF(ROUND(D%d,0)=0,"OK","DIFFERENZA")' % n)
        n = v.max_row + 2
        v.cell(n, 1, "data valuta chiude come originale")
        v.cell(n, 2, "=originale!F%d" % last1)
        v.cell(n, 3, "='data valuta'!F%d" % last2)
        v.cell(n, 4, "=B%d-C%d" % (n, n))
        v.cell(n, 5, '=IF(ROUND(D%d,0)=0,"OK","DIFFERENZA")' % n)
        v.cell(n + 2, 1, "Foglio cancellabile: non serve al ricalcolo. "
                         "Aggiungere qui le mappature descrizione e i controlli assegni.")
        for col, w in zip("ABCDE", (14, 18, 18, 14, 14)):
            v.column_dimensions[col].width = w

    wb.save(a.out)
    print("%s  originale rows->%d  data valuta rows->%d  %d statements"
          % (a.out, last1, last2, len(sts)))
    print("now run: python /mnt/skills/public/xlsx/scripts/recalc.py %s" % a.out)


if __name__ == "__main__":
    main()
